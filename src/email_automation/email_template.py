from docx import Document

TEMPLATE = "src/email_automation/templates/Violation_Email_Template.docx"

def create_letter(row):
    doc = Document(TEMPLATE)

    replacements = {
        "{SKU}": str(row["SKU"]),
        "{Seller_Name}": str(row["Seller_name"]),
        "{PL}": str(row["PL"]),
        "{Marketplace}": str(row["Marketplace"]),
        "{Violation_date}": str(row["Violation_date"]),
        "{Advertised_price}": str(row["Advertised_price"]),
        "{Promotional_price}": str(row["Promotional_price"]),
        "{Violation_count}": str(row["Violation_count"]),
        "{Action}": str(row["Action"])
    }

    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            paragraph.text = paragraph.text.replace(old, new)

    for table in doc.tables:
        for row_cells in table.rows:
            for cell in row_cells.cells:
                for old, new in replacements.items():
                    cell.text = cell.text.replace(old, new)

    filename = f"Violation_{row['SKU']}_{row['Violation_count']}.docx"
    doc.save(filename)

    return filename